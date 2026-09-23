"""Lossless XLSX container for the versioned questions-state envelope.

The JSON state envelope remains the canonical wire model.  This module only
changes its container: each entity type gets a worksheet and the reader
reconstructs the exact same dictionary that the JSON importer validates.

Images are split into 30,000-character base64 chunks because an Excel cell is
limited to 32,767 characters.  Text cells are explicitly marked as strings so
content beginning with ``=``, ``+``, ``-`` or ``@`` is never emitted as a
formula.
"""

import json
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


STATE_WORKBOOK_FORMAT = 'mukhtabir-state-xlsx'
STATE_WORKBOOK_VERSION = 1
IMAGE_CHUNK_SIZE = 30_000

_SHEET_COLUMNS = {
    'Metadata': ('key', 'value'),
    'Questions': (
        'uuid', 'question', 'choices_json', 'correct_answer', 'explanation',
        'source', 'difficulty', 'category_uuid', 'case_uuid', 'case_order',
        'is_draft', 'verified', 'verified_by',
        'verified_at', 'verification_notes', 'authored_by_uuid',
        'authored_by_name', 'owned_by_uuid', 'owned_by_name', 'updated_by',
        'times_answered', 'times_correct', 'version', 'created_at',
        'updated_at',
    ),
    'Categories': (
        'uuid', 'name', 'description', 'color', 'icon',
    ),
    'Tags': ('uuid', 'name', 'parent_uuid'),
    'QuestionTags': ('question_uuid', 'tag_uuid'),
    'Cases': (
        'uuid', 'key', 'title', 'stem', 'authored_by_uuid',
        'authored_by_name',
    ),
    'Users': ('uuid', 'username', 'full_name'),
    'Images': (
        'question_uuid', 'filename', 'mime', 'chunk_index',
        'data_base64_chunk',
    ),
}


class StateWorkbookError(ValueError):
    """Raised when an XLSX cannot be decoded as a state workbook."""


def _write_cell(sheet, value):
    """Return a write-only cell, forcing all Python strings to literals."""
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def _append_row(sheet, values):
    sheet.append([_write_cell(sheet, value) for value in values])


def _text(value, *, nullable=False):
    if value is None or value == '':
        return None if nullable else ''
    return str(value)


def _boolean(value, field):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().lower()
    if normalized in {'true', '1', 'yes'}:
        return True
    if normalized in {'false', '0', 'no'}:
        return False
    raise StateWorkbookError(f'Invalid boolean in {field}')


def _integer(value, field, *, nullable=False, default=None):
    if value is None or value == '':
        if nullable:
            return None
        if default is not None:
            return default
        raise StateWorkbookError(f'Missing integer in {field}')
    if isinstance(value, bool):
        raise StateWorkbookError(f'Invalid integer in {field}')
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError
        return int(number)
    except (InvalidOperation, TypeError, ValueError, OverflowError):
        raise StateWorkbookError(f'Invalid integer in {field}') from None


def _json_list(value, field):
    try:
        parsed = json.loads(_text(value))
    except (TypeError, ValueError, json.JSONDecodeError):
        raise StateWorkbookError(f'Invalid JSON list in {field}') from None
    if not isinstance(parsed, list):
        raise StateWorkbookError(f'Invalid JSON list in {field}')
    return parsed


def _worksheet_rows(workbook, name):
    """Yield dictionaries from a worksheet while validating its header."""
    if name not in workbook.sheetnames:
        raise StateWorkbookError(f'Missing worksheet: {name}')

    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration:
        raise StateWorkbookError(f'Worksheet {name} is empty') from None

    header = [_text(value).strip() for value in raw_header]
    if any(not value for value in header) or len(header) != len(set(header)):
        raise StateWorkbookError(f'Worksheet {name} has an invalid header')

    missing = set(_SHEET_COLUMNS[name]) - set(header)
    if missing:
        raise StateWorkbookError(
            f'Worksheet {name} is missing columns: {", ".join(sorted(missing))}'
        )

    for raw_row in rows:
        values = list(raw_row) + [None] * max(0, len(header) - len(raw_row))
        row = dict(zip(header, values))
        if not any(value is not None and value != '' for value in row.values()):
            continue
        yield row


def _validate_zip_size(path, max_uncompressed_size):
    """Reject obvious XLSX zip bombs before openpyxl expands their XML."""
    if max_uncompressed_size is None:
        return
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 1_000:
                raise StateWorkbookError('Workbook contains too many files')
            if sum(member.file_size for member in members) > max_uncompressed_size:
                raise StateWorkbookError('Workbook expands beyond the import limit')
    except BadZipFile:
        raise StateWorkbookError('Invalid XLSX file') from None


def write_state_workbook(payload, filepath):
    """Serialize a validated state-envelope dictionary as a multi-sheet XLSX."""
    workbook = Workbook(write_only=True)

    metadata = workbook.create_sheet('Metadata')
    _append_row(metadata, _SHEET_COLUMNS['Metadata'])
    meta = payload.get('meta') or {}
    metadata_rows = (
        ('workbook_format', STATE_WORKBOOK_FORMAT),
        ('workbook_version', STATE_WORKBOOK_VERSION),
        ('state_format', meta.get('format') or ''),
        ('state_version', meta.get('version')),
        ('exported_at', meta.get('exported_at') or ''),
        ('includes_images', bool(meta.get('includes_images'))),
    )
    for row in metadata_rows:
        _append_row(metadata, row)

    questions = workbook.create_sheet('Questions')
    _append_row(questions, _SHEET_COLUMNS['Questions'])
    for item in payload.get('questions') or []:
        _append_row(questions, (
            item.get('uuid'), item.get('question'),
            json.dumps(item.get('choices') or [], ensure_ascii=False),
            item.get('correct_answer'), item.get('explanation') or '',
            item.get('source') or '', item.get('difficulty') or 'medium',
            item.get('category_uuid') or '',
            item.get('case_uuid') or '', item.get('case_order'),
            bool(item.get('is_draft')), bool(item.get('verified')),
            item.get('verified_by') or '', item.get('verified_at') or '',
            item.get('verification_notes') or '',
            item.get('authored_by_uuid') or '',
            item.get('authored_by_name') or '',
            item.get('owned_by_uuid') or '', item.get('owned_by_name') or '',
            item.get('updated_by') or '', item.get('times_answered', 0),
            item.get('times_correct', 0), item.get('version', 1),
            item.get('created_at') or '', item.get('updated_at') or '',
        ))

    categories = workbook.create_sheet('Categories')
    _append_row(categories, _SHEET_COLUMNS['Categories'])
    for item in payload.get('categories') or []:
        _append_row(categories, (
            item.get('uuid'), item.get('name'), item.get('description') or '',
            item.get('color') or '', item.get('icon') or '',
        ))

    tags = workbook.create_sheet('Tags')
    _append_row(tags, _SHEET_COLUMNS['Tags'])
    for item in payload.get('tags') or []:
        _append_row(tags, (
            item.get('uuid'), item.get('name'), item.get('parent_uuid') or '',
        ))

    question_tags = workbook.create_sheet('QuestionTags')
    _append_row(question_tags, _SHEET_COLUMNS['QuestionTags'])
    for question in payload.get('questions') or []:
        for tag_uuid in question.get('tags') or []:
            _append_row(question_tags, (question.get('uuid'), tag_uuid))

    cases = workbook.create_sheet('Cases')
    _append_row(cases, _SHEET_COLUMNS['Cases'])
    for item in payload.get('cases') or []:
        _append_row(cases, (
            item.get('uuid'), item.get('key'), item.get('title') or '',
            item.get('stem') or '', item.get('authored_by_uuid') or '',
            item.get('authored_by_name') or '',
        ))

    users = workbook.create_sheet('Users')
    _append_row(users, _SHEET_COLUMNS['Users'])
    for user_uuid, item in (meta.get('user_map') or {}).items():
        _append_row(users, (
            user_uuid, item.get('username') or '', item.get('full_name') or '',
        ))

    images = workbook.create_sheet('Images')
    _append_row(images, _SHEET_COLUMNS['Images'])
    for question in payload.get('questions') or []:
        image = question.get('image')
        if not isinstance(image, dict) or not image.get('data_base64'):
            continue
        encoded = image['data_base64']
        for offset in range(0, len(encoded), IMAGE_CHUNK_SIZE):
            _append_row(images, (
                question.get('uuid'), image.get('filename') or '',
                image.get('mime') or '', offset // IMAGE_CHUNK_SIZE + 1,
                encoded[offset:offset + IMAGE_CHUNK_SIZE],
            ))

    workbook.save(Path(filepath))


def read_state_workbook(filepath, *, max_uncompressed_size=None):
    """Read a state XLSX and return the canonical state-envelope dictionary."""
    _validate_zip_size(filepath, max_uncompressed_size)
    try:
        workbook = load_workbook(filepath, read_only=True, data_only=False)
    except Exception as exc:
        raise StateWorkbookError('Invalid XLSX file') from exc

    try:
        metadata = {}
        for row in _worksheet_rows(workbook, 'Metadata'):
            key = _text(row.get('key')).strip()
            if not key or key in metadata:
                raise StateWorkbookError('Metadata contains a duplicate or empty key')
            metadata[key] = row.get('value')

        if metadata.get('workbook_format') != STATE_WORKBOOK_FORMAT:
            raise StateWorkbookError('Not a Mukhtabir state workbook')
        if _integer(metadata.get('workbook_version'), 'workbook_version') != STATE_WORKBOOK_VERSION:
            raise StateWorkbookError('Unsupported state workbook version')

        payload = {
            'meta': {
                'format': _text(metadata.get('state_format')),
                'version': _integer(metadata.get('state_version'), 'state_version'),
                'exported_at': _text(metadata.get('exported_at')),
                'includes_images': _boolean(
                    metadata.get('includes_images'), 'includes_images',
                ),
                'user_map': {},
                'counts': {},
            },
            'categories': [],
            'tags': [],
            'cases': [],
            'questions': [],
        }

        for row in _worksheet_rows(workbook, 'Categories'):
            payload['categories'].append({
                'uuid': _text(row['uuid']),
                'name': _text(row['name']),
                'description': _text(row['description']),
                'color': _text(row['color']),
                'icon': _text(row['icon']),
            })

        for row in _worksheet_rows(workbook, 'Tags'):
            payload['tags'].append({
                'uuid': _text(row['uuid']),
                'name': _text(row['name']),
                'parent_uuid': _text(row['parent_uuid'], nullable=True),
            })

        for row in _worksheet_rows(workbook, 'Cases'):
            payload['cases'].append({
                'uuid': _text(row['uuid']),
                'key': _text(row['key']),
                'title': _text(row['title']),
                'stem': _text(row['stem']),
                'authored_by_uuid': _text(row['authored_by_uuid'], nullable=True),
                'authored_by_name': _text(row['authored_by_name'], nullable=True),
            })

        for row in _worksheet_rows(workbook, 'Users'):
            user_uuid = _text(row['uuid'])
            if user_uuid in payload['meta']['user_map']:
                raise StateWorkbookError('Users contains a duplicate uuid')
            payload['meta']['user_map'][user_uuid] = {
                'username': _text(row['username']),
                'full_name': _text(row['full_name']),
            }

        question_by_uuid = {}
        for row in _worksheet_rows(workbook, 'Questions'):
            question = {
                'uuid': _text(row['uuid']),
                'question': _text(row['question']),
                'choices': _json_list(row['choices_json'], 'choices_json'),
                'correct_answer': _integer(row['correct_answer'], 'correct_answer'),
                'explanation': _text(row['explanation']),
                'source': _text(row['source']),
                'difficulty': _text(row['difficulty']),
                'category_uuid': _text(row['category_uuid'], nullable=True),
                'tags': [],
                'case_uuid': _text(row['case_uuid'], nullable=True),
                'case_order': _integer(
                    row['case_order'], 'case_order', nullable=True,
                ),
                'is_draft': _boolean(row['is_draft'], 'is_draft'),
                'verified': _boolean(row['verified'], 'verified'),
                'verified_by': _text(row['verified_by'], nullable=True),
                'verified_at': _text(row['verified_at'], nullable=True),
                'verification_notes': _text(
                    row['verification_notes'], nullable=True,
                ),
                'authored_by_uuid': _text(
                    row['authored_by_uuid'], nullable=True,
                ),
                'authored_by_name': _text(
                    row['authored_by_name'], nullable=True,
                ),
                'owned_by_uuid': _text(row['owned_by_uuid'], nullable=True),
                'owned_by_name': _text(row['owned_by_name'], nullable=True),
                'updated_by': _text(row['updated_by'], nullable=True),
                'times_answered': _integer(
                    row['times_answered'], 'times_answered', default=0,
                ),
                'times_correct': _integer(
                    row['times_correct'], 'times_correct', default=0,
                ),
                'version': _integer(row['version'], 'version', default=1),
                'created_at': _text(row['created_at'], nullable=True),
                'updated_at': _text(row['updated_at'], nullable=True),
                'image': None,
            }
            if question['uuid'] in question_by_uuid:
                raise StateWorkbookError('Questions contains a duplicate uuid')
            question_by_uuid[question['uuid']] = question
            payload['questions'].append(question)

        for row in _worksheet_rows(workbook, 'QuestionTags'):
            question_uuid = _text(row['question_uuid'])
            if question_uuid not in question_by_uuid:
                raise StateWorkbookError(
                    'QuestionTags references an unknown question'
                )
            question_by_uuid[question_uuid]['tags'].append(
                _text(row['tag_uuid'])
            )

        image_parts = defaultdict(list)
        image_meta = {}
        for row in _worksheet_rows(workbook, 'Images'):
            question_uuid = _text(row['question_uuid'])
            if question_uuid not in question_by_uuid:
                raise StateWorkbookError('Images references an unknown question')
            filename = _text(row['filename'])
            mime = _text(row['mime'])
            current_meta = (filename, mime)
            if question_uuid in image_meta and image_meta[question_uuid] != current_meta:
                raise StateWorkbookError('Image chunks have inconsistent metadata')
            image_meta[question_uuid] = current_meta
            image_parts[question_uuid].append((
                _integer(row['chunk_index'], 'chunk_index'),
                _text(row['data_base64_chunk']),
            ))

        for question_uuid, parts in image_parts.items():
            parts.sort(key=lambda item: item[0])
            expected = list(range(1, len(parts) + 1))
            if [item[0] for item in parts] != expected:
                raise StateWorkbookError('Image chunks are missing or duplicated')
            filename, mime = image_meta[question_uuid]
            question_by_uuid[question_uuid]['image'] = {
                'filename': filename,
                'mime': mime,
                'data_base64': ''.join(item[1] for item in parts),
            }

        payload['meta']['counts'] = {
            'categories': len(payload['categories']),
            'tags': len(payload['tags']),
            'cases': len(payload['cases']),
            'questions': len(payload['questions']),
            'users': len(payload['meta']['user_map']),
        }
        return payload
    finally:
        workbook.close()
