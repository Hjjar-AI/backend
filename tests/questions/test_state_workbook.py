import tempfile
from pathlib import Path

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.questions.services.state_workbook import (
    IMAGE_CHUNK_SIZE,
    StateWorkbookError,
    read_state_workbook,
    write_state_workbook,
)


QUESTION_UUID = '33333333-3333-4333-8333-333333333333'
KNOWLEDGE_UUID = '55555555-5555-4555-8555-555555555555'


def _payload():
    return {
        'meta': {
            'format': 'mukhtabir-questions',
            'version': 4,
            'exported_at': '2026-01-01T00:00:00+00:00',
            'includes_images': True,
            'scope': 'selection',
            'selection': {'difficulty': 'easy'},
            'user_map': {
                '11111111-1111-4111-8111-111111111111': {
                    'username': '=literal-author',
                    'full_name': 'Literal Author',
                },
            },
            'counts': {},
        },
        'categories': [{
            'uuid': '22222222-2222-4222-8222-222222222222',
            'name': '=literal-category',
            'description': '',
            'color': '#112233',
            'icon': 'bi-folder',
        }],
        'tags': [{
            'uuid': '44444444-4444-4444-8444-444444444444',
            'name': 'tag',
            'parent_uuid': None,
        }],
        'cases': [],
        'knowledge_objects': [{
            'uuid': KNOWLEDGE_UUID,
            'title': 'Literal concept',
            'learning_objective': 'Explain the concept',
            'canonical_answer': 'Canonical answer',
            'key_facts': ['Fact one'],
            'misconceptions': ['Common misconception'],
            'category_uuid': '22222222-2222-4222-8222-222222222222',
            'tags': ['44444444-4444-4444-8444-444444444444'],
            'source_document': 'source.pdf',
            'source_page': 12,
            'translations': {'ar': {'title': 'مفهوم'}},
            'status': 'active',
            'version': 2,
            'last_revised_at': '2026-01-02',
            'created_by_uuid': '11111111-1111-4111-8111-111111111111',
        }],
        'questions': [{
            'uuid': QUESTION_UUID,
            'question': '=literal question',
            'choices': ['+literal choice', 'Other'],
            'correct_answer': 1,
            'explanation': '',
            'source': '@literal-source',
            'source_document': 'source.pdf',
            'source_page': 12,
            'translations': {
                'ar': {
                    'question': 'سؤال',
                    'choices': ['أ', 'ب'],
                    'explanation': 'شرح',
                },
            },
            'difficulty': 'easy',
            'knowledge_object_uuid': KNOWLEDGE_UUID,
            'last_revised_at': '2026-01-02',
            'category_uuid': '22222222-2222-4222-8222-222222222222',
            'tags': ['44444444-4444-4444-8444-444444444444'],
            'case_uuid': None,
            'case_order': None,
            'is_draft': False,
            'verified': True,
            'verified_by': 'reviewer',
            'verified_at': '2026-01-02T00:00:00+00:00',
            'verification_notes': None,
            'authored_by_uuid': '11111111-1111-4111-8111-111111111111',
            'authored_by_name': '=literal-author',
            'owned_by_uuid': None,
            'owned_by_name': None,
            'updated_by': '@auditor',
            'times_answered': 7,
            'times_correct': 5,
            'version': 4,
            'created_at': '2025-01-01T00:00:00+00:00',
            'updated_at': '2026-01-03T00:00:00+00:00',
            'image': {
                'filename': 'question.png',
                'mime': 'image/png',
                'data_base64': 'A' * (IMAGE_CHUNK_SIZE * 2 + 17),
            },
        }],
    }


class StateWorkbookTests(SimpleTestCase):
    def test_round_trip_preserves_all_sections_stats_and_chunked_image(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'state.xlsx'
            write_state_workbook(_payload(), path)
            restored = read_state_workbook(
                path, max_uncompressed_size=10 * 1024 * 1024,
            )

        question = restored['questions'][0]
        self.assertEqual(question['question'], '=literal question')
        self.assertEqual(question['choices'], ['+literal choice', 'Other'])
        self.assertEqual(question['times_answered'], 7)
        self.assertEqual(question['times_correct'], 5)
        self.assertEqual(question['version'], 4)
        self.assertEqual(question['source_document'], 'source.pdf')
        self.assertEqual(question['source_page'], 12)
        self.assertEqual(question['translations'], _payload()['questions'][0]['translations'])
        self.assertEqual(question['knowledge_object_uuid'], KNOWLEDGE_UUID)
        self.assertEqual(question['last_revised_at'], '2026-01-02')
        self.assertEqual(
            restored['knowledge_objects'], _payload()['knowledge_objects'],
        )
        self.assertEqual(restored['meta']['scope'], 'selection')
        self.assertEqual(restored['meta']['selection'], {'difficulty': 'easy'})
        self.assertEqual(
            question['tags'], ['44444444-4444-4444-8444-444444444444'],
        )
        self.assertEqual(
            question['image']['data_base64'],
            _payload()['questions'][0]['image']['data_base64'],
        )
        self.assertEqual(restored['categories'], _payload()['categories'])
        self.assertEqual(
            restored['meta']['user_map'], _payload()['meta']['user_map'],
        )

    def test_formula_like_text_is_stored_as_literal_string(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'state.xlsx'
            write_state_workbook(_payload(), path)
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                cell = workbook['Questions']['B2']
                self.assertEqual(cell.value, '=literal question')
                self.assertEqual(cell.data_type, 's')
            finally:
                workbook.close()

    def test_non_state_workbook_is_rejected(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'ordinary.xlsx'
            workbook = Workbook()
            workbook.save(path)

            with self.assertRaises(StateWorkbookError):
                read_state_workbook(path)

    def test_legacy_v2_workbook_keeps_new_question_fields_absent(self):
        """Migration can distinguish "old field absent" from "clear it"."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'legacy-state.xlsx'
            write_state_workbook(_payload(), path)
            workbook = load_workbook(path)
            try:
                metadata = workbook['Metadata']
                for row in metadata.iter_rows(min_row=2):
                    if row[0].value == 'state_version':
                        row[1].value = 2
                questions = workbook['Questions']
                header = [cell.value for cell in questions[1]]
                for name in (
                    'translations_json', 'source_page', 'source_document',
                    'knowledge_object_uuid', 'last_revised_at',
                ):
                    questions.delete_cols(header.index(name) + 1)
                    header.remove(name)
                del workbook['KnowledgeObjects']
                workbook.save(path)
            finally:
                workbook.close()

            restored = read_state_workbook(path)

        question = restored['questions'][0]
        self.assertNotIn('source_document', question)
        self.assertNotIn('source_page', question)
        self.assertNotIn('translations', question)
        self.assertNotIn('knowledge_object_uuid', question)
        self.assertNotIn('last_revised_at', question)
        self.assertEqual(restored['knowledge_objects'], [])
